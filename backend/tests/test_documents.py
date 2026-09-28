"""
test_documents.py — Verifies document generation tools and structural validity.
Tests:
1. Word (.docx) document generation with title, sections, paragraphs.
2. Excel (.xlsx) workbook generation with styled headers and data rows.
3. PowerPoint (.pptx) presentation with slides and bullet points.
4. Python code (.py) script generation.
"""
from __future__ import annotations

import sys
from pathlib import Path

backend_dir = Path(__file__).resolve().parent.parent
if str(backend_dir) not in sys.path:
    sys.path.insert(0, str(backend_dir))

from tools.documents import (
    generate_code_file,
    generate_excel_sheet,
    generate_powerpoint,
    generate_word_doc,
)


def test_document_tools():
    print("=== Testing Document Generation Tools ===")

    # 1. Word Document
    word_content = {
        "title": "Artificial Intelligence Overview",
        "sections": [
            {
                "heading": "Introduction",
                "paragraphs": [
                    "Artificial intelligence is transforming automation and developer workflows.",
                    "Agentic AI systems leverage LLMs for multi-step reasoning.",
                ],
            }
        ],
    }
    word_path = Path(generate_word_doc("test_ai_doc.docx", word_content))
    assert word_path.exists(), "Word document was not created"
    assert word_path.stat().st_size > 1000, "Word document is unexpectedly small"
    from docx import Document

    doc = Document(str(word_path))
    assert len(doc.paragraphs) >= 2
    print(f"  [OK] Word document verified ({len(doc.paragraphs)} paragraphs, {word_path.stat().st_size} bytes)")

    # 2. Excel Spreadsheet
    excel_content = {
        "sheet_name": "Students",
        "headers": ["Roll No", "Student Name", "Attendance %"],
        "rows": [
            [101, "Alice Smith", "96%"],
            [102, "Bob Jones", "92%"],
            [103, "Charlie Brown", "98%"],
        ],
    }
    excel_path = Path(generate_excel_sheet("test_attendance.xlsx", excel_content))
    assert excel_path.exists(), "Excel workbook was not created"
    import openpyxl

    wb = openpyxl.load_workbook(str(excel_path))
    sheet = wb.active
    assert sheet is not None and sheet.max_row == 4
    print(f"  [OK] Excel workbook verified ('{sheet.title}', {sheet.max_row} rows, {sheet.max_column} columns)")

    # 3. PowerPoint Presentation
    ppt_content = {
        "title": "Cloud Computing Fundamentals",
        "subtitle": "Infrastructure as a Service & Beyond",
        "slides": [
            {"title": "Overview", "bullets": ["Scalability", "High Availability", "Cost Efficiency"]},
            {"title": "Architecture", "bullets": ["Microservices", "Serverless", "Storage"]},
        ],
    }
    ppt_path = Path(generate_powerpoint("test_cloud.pptx", ppt_content))
    assert ppt_path.exists(), "PowerPoint presentation was not created"
    from pptx import Presentation

    prs = Presentation(str(ppt_path))
    assert len(prs.slides) == 3  # title + 2 content slides
    print(f"  [OK] PowerPoint presentation verified ({len(prs.slides)} slides)")

    # 4. Code Generation
    code_content = {"code": "def fib(n):\n    return n if n <= 1 else fib(n-1) + fib(n-2)\n"}
    code_path = Path(generate_code_file("test_fib.py", code_content))
    assert code_path.exists(), "Code file was not created"
    content = code_path.read_text(encoding="utf-8")
    assert "def fib(n):" in content
    print(f"  [OK] Code generation verified ({len(content)} characters)")

    print("[SUCCESS] All document generation tools verified successfully!\n")


if __name__ == "__main__":
    test_document_tools()
