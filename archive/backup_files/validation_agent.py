"""
validation_agent.py — Validates output of plan execution and triggers replan/reask on failure.
"""
from __future__ import annotations

import logging
import os
from typing import Any

from graph.state import AgentState

logger = logging.getLogger(__name__)

async def validation_agent_node(state: AgentState) -> dict[str, Any]:
    """LangGraph node: validates execution results."""
    error = state.get("error")
    last_res = state.get("last_execution_result") or {}
    plan = state.get("plan") or []

    failure_msg = None

    # 1. Check for execution errors
    if error or last_res.get("success") is False:
        failure_msg = error or last_res.get("error", "Unknown execution failure")
    
    # 2. Check for generated files and validate them strictly
    elif "output_file" in last_res:
        filepath = last_res["output_file"]
        if not os.path.exists(filepath):
            failure_msg = f"Output file does not exist: {filepath}"
        elif os.path.getsize(filepath) == 0:
            failure_msg = f"Output file is completely empty (0 bytes): {filepath}"
        else:
            # Type specific internal checks
            ext = os.path.splitext(filepath)[1].lower()
            try:
                if ext == ".pptx":
                    import pptx
                    prs = pptx.Presentation(filepath)
                    if len(prs.slides) == 0:
                        failure_msg = "PowerPoint generated successfully but contains no slides."
                    else:
                        has_text = False
                        for slide in prs.slides:
                            for shape in slide.shapes:
                                if hasattr(shape, "text") and shape.text.strip():
                                    has_text = True
                                    break
                            if has_text:
                                break
                        if not has_text:
                            failure_msg = "PowerPoint was generated but all slides are completely empty with no text."
                            
                elif ext == ".docx":
                    import docx
                    doc = docx.Document(filepath)
                    if len(doc.paragraphs) == 0:
                        failure_msg = "Word document generated successfully but contains no paragraphs."
                    else:
                        has_text = any(p.text.strip() for p in doc.paragraphs)
                        if not has_text:
                            failure_msg = "Word document generated successfully but has no visible text content."
            except Exception as e:
                failure_msg = f"File created but failed internal structure validation: {e}"

    if failure_msg:
        logger.warning("Validation Agent detected failure: %s", failure_msg)
        logger.info("VALIDATION RESULT: Failed - %s", failure_msg)
        return {
            "validation_result": {"valid": False, "reason": failure_msg},
            "replan_reason": failure_msg,
            "execution_trace": [
                {
                    "agent": "validation_agent",
                    "status": "error",
                    "message": f"Validation failed: {failure_msg}",
                }
            ],
        }

    # Validation passed
    success_msg = f"Validated execution of {len(plan)} plan steps successfully."
    logger.info("VALIDATION RESULT: Success - %s", success_msg)
    return {
        "validation_result": {"valid": True, "reason": success_msg},
        "replan_reason": None,
        "execution_trace": [
            {
                "agent": "validation_agent",
                "status": "success",
                "message": success_msg,
            }
        ],
    }
